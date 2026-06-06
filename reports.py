import os
import logging
from datetime import date
from typing import List, Dict

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

logger = logging.getLogger(__name__)

UZ_MONTHS = {
    1: "yanvar", 2: "fevral", 3: "mart", 4: "aprel",
    5: "may",    6: "iyun",   7: "iyul",  8: "avgust",
    9: "sentabr", 10: "oktabr", 11: "noyabr", 12: "dekabr",
}

_THIN = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center")
_LEFT   = Alignment(horizontal="left", vertical="center", indent=1)


def _hdr_cell(ws, row, col, value, bg="2E4057"):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(bold=True, color="FFFFFF", size=11)
    c.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
    c.alignment = _CENTER
    c.border = _BORDER
    return c


def generate_monthly_excel(
    incomes: List[Dict], year: int, month: int, output_dir: str = "/tmp"
) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{UZ_MONTHS[month][:3]} {year}"

    # ── Title row ──────────────────────────────────────────────────────────
    ws.merge_cells("A1:G1")
    t = ws["A1"]
    t.value = f"DAROMAD HISOBOTI  —  {UZ_MONTHS[month].upper()} {year}"
    t.font = Font(bold=True, size=13)
    t.alignment = _CENTER
    ws.row_dimensions[1].height = 28

    # ── Column headers ─────────────────────────────────────────────────────
    headers = ["#", "Menejer", "Tur", "Summa", "Valyuta", "USD", "Sana"]
    col_widths = [5, 22, 28, 14, 10, 12, 14]
    for col, (h, w) in enumerate(zip(headers, col_widths), 1):
        _hdr_cell(ws, 2, col, h)
        ws.column_dimensions[ws.cell(row=2, column=col).column_letter].width = w
    ws.row_dimensions[2].height = 20

    # ── Data rows ──────────────────────────────────────────────────────────
    alt_fill     = PatternFill(start_color="F4F6F8", end_color="F4F6F8", fill_type="solid")
    worker_fill  = PatternFill(start_color="E3F2E8", end_color="E3F2E8", fill_type="solid")
    sub_fill     = PatternFill(start_color="D0E9F7", end_color="D0E9F7", fill_type="solid")
    grand_fill   = PatternFill(start_color="2E4057", end_color="2E4057", fill_type="solid")

    by_worker: dict = {}
    for inc in incomes:
        name = inc.get("full_name") or inc.get("username") or str(inc["user_id"])
        by_worker.setdefault(name, []).append(inc)

    row = 3
    grand_total = 0.0

    for worker_name, recs in by_worker.items():
        # Worker sub-header
        ws.merge_cells(f"A{row}:G{row}")
        c = ws.cell(row=row, column=1, value=f"  {worker_name}")
        c.font = Font(bold=True, size=11)
        c.fill = worker_fill
        c.alignment = _LEFT
        ws.row_dimensions[row].height = 18
        row += 1

        subtotal = 0.0
        for idx, inc in enumerate(recs):
            bg = alt_fill if idx % 2 == 0 else None
            cells_data = [
                idx + 1,
                worker_name,
                inc["description"],
                inc["amount"],
                inc["currency"],
                round(inc["amount_usd"], 2),
                inc["created_at"][:10],
            ]
            for col, val in enumerate(cells_data, 1):
                c = ws.cell(row=row, column=col, value=val)
                c.border = _BORDER
                c.alignment = _CENTER if col != 3 else _LEFT
                if bg:
                    c.fill = bg
                if col == 4:
                    c.number_format = "#,##0.00"
                if col == 6:
                    c.number_format = '"$"#,##0.00'
            subtotal += inc["amount_usd"]
            row += 1

        # Subtotal row
        ws.merge_cells(f"A{row}:E{row}")
        lbl = ws.cell(row=row, column=1, value=f"  {worker_name} jami:")
        lbl.font = Font(bold=True)
        lbl.fill = sub_fill
        lbl.alignment = _LEFT
        total_c = ws.cell(row=row, column=6, value=round(subtotal, 2))
        total_c.font = Font(bold=True)
        total_c.fill = sub_fill
        total_c.number_format = '"$"#,##0.00'
        total_c.alignment = _CENTER
        ws.cell(row=row, column=7).fill = sub_fill
        for col in range(1, 8):
            ws.cell(row=row, column=col).border = _BORDER
        row += 1
        grand_total += subtotal

    # ── Grand total ────────────────────────────────────────────────────────
    row += 1
    ws.merge_cells(f"A{row}:E{row}")
    g_lbl = ws.cell(row=row, column=1, value="  UMUMIY JAMI")
    g_lbl.font = Font(bold=True, color="FFFFFF", size=12)
    g_lbl.fill = grand_fill
    g_lbl.alignment = _LEFT

    g_val = ws.cell(row=row, column=6, value=round(grand_total, 2))
    g_val.font = Font(bold=True, color="FFFFFF", size=12)
    g_val.fill = grand_fill
    g_val.number_format = '"$"#,##0.00'
    g_val.alignment = _CENTER

    g_cur = ws.cell(row=row, column=7, value="USD")
    g_cur.font = Font(bold=True, color="FFFFFF")
    g_cur.fill = grand_fill
    g_cur.alignment = _CENTER

    for col in range(1, 8):
        ws.cell(row=row, column=col).border = _BORDER
    ws.row_dimensions[row].height = 24

    ws.freeze_panes = "A3"

    os.makedirs(output_dir, exist_ok=True)
    filepath = os.path.join(output_dir, f"hisobot_{year}_{month:02d}.xlsx")
    wb.save(filepath)
    logger.info(f"Excel report saved: {filepath}")
    return filepath


def generate_worker_excel(
    incomes: List[Dict], worker_name: str, year: int, month: int, output_dir: str = "/tmp"
) -> str:
    """Shaxsiy Excel hisobot — hafta bo'yicha breakdown."""
    from datetime import date as dt, timedelta

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{UZ_MONTHS[month][:3]} {year}"

    # Title
    ws.merge_cells("A1:F1")
    t = ws["A1"]
    t.value = f"{worker_name.upper()} — {UZ_MONTHS[month].upper()} {year}"
    t.font = Font(bold=True, size=13)
    t.alignment = _CENTER
    ws.row_dimensions[1].height = 28

    # Headers
    headers = ["#", "Tur", "Summa", "Valyuta", "USD (~)", "Sana"]
    col_widths = [5, 30, 14, 10, 12, 14]
    for col, (h, w) in enumerate(zip(headers, col_widths), 1):
        _hdr_cell(ws, 2, col, h)
        ws.column_dimensions[ws.cell(row=2, column=col).column_letter].width = w
    ws.row_dimensions[2].height = 20

    alt_fill   = PatternFill(start_color="F4F6F8", end_color="F4F6F8", fill_type="solid")
    week_fill  = PatternFill(start_color="E3F2E8", end_color="E3F2E8", fill_type="solid")
    sub_fill   = PatternFill(start_color="D0E9F7", end_color="D0E9F7", fill_type="solid")
    grand_fill = PatternFill(start_color="2E4057", end_color="2E4057", fill_type="solid")

    # Group by week_start
    by_week: dict = {}
    for inc in incomes:
        by_week.setdefault(inc["week_start"], []).append(inc)

    row = 3
    grand_total = 0.0

    for week_str, recs in sorted(by_week.items()):
        ws_date = dt.fromisoformat(week_str)
        we_date = ws_date + timedelta(days=6)
        if ws_date.month == we_date.month:
            wrange = f"{ws_date.day}–{we_date.day} {UZ_MONTHS[ws_date.month]}"
        else:
            wrange = f"{ws_date.day} {UZ_MONTHS[ws_date.month]} – {we_date.day} {UZ_MONTHS[we_date.month]}"

        # Week header row
        ws.merge_cells(f"A{row}:F{row}")
        c = ws.cell(row=row, column=1, value=f"  📅 {wrange} hafta")
        c.font = Font(bold=True)
        c.fill = week_fill
        c.alignment = _LEFT
        ws.row_dimensions[row].height = 18
        row += 1

        subtotal = 0.0
        for idx, inc in enumerate(recs):
            bg = alt_fill if idx % 2 == 0 else None
            data = [idx + 1, inc["description"], inc["amount"],
                    inc["currency"], round(inc["amount_usd"], 2), inc["created_at"][:10]]
            for col, val in enumerate(data, 1):
                c = ws.cell(row=row, column=col, value=val)
                c.border = _BORDER
                c.alignment = _CENTER if col != 2 else _LEFT
                if bg:
                    c.fill = bg
                if col == 3:
                    c.number_format = "#,##0.00"
                if col == 5:
                    c.number_format = '"$"#,##0.00'
            subtotal += inc["amount_usd"]
            row += 1

        # Week subtotal
        ws.merge_cells(f"A{row}:D{row}")
        lbl = ws.cell(row=row, column=1, value="  Hafta jami:")
        lbl.font = Font(bold=True)
        lbl.fill = sub_fill
        lbl.alignment = _LEFT
        tv = ws.cell(row=row, column=5, value=round(subtotal, 2))
        tv.font = Font(bold=True)
        tv.fill = sub_fill
        tv.number_format = '"$"#,##0.00'
        tv.alignment = _CENTER
        ws.cell(row=row, column=6).fill = sub_fill
        for col in range(1, 7):
            ws.cell(row=row, column=col).border = _BORDER
        row += 1
        grand_total += subtotal

    # Grand total
    row += 1
    ws.merge_cells(f"A{row}:D{row}")
    g_lbl = ws.cell(row=row, column=1, value="  UMUMIY JAMI")
    g_lbl.font = Font(bold=True, color="FFFFFF", size=12)
    g_lbl.fill = grand_fill
    g_lbl.alignment = _LEFT
    g_val = ws.cell(row=row, column=5, value=round(grand_total, 2))
    g_val.font = Font(bold=True, color="FFFFFF", size=12)
    g_val.fill = grand_fill
    g_val.number_format = '"$"#,##0.00'
    g_val.alignment = _CENTER
    g_cur = ws.cell(row=row, column=6, value="USD")
    g_cur.font = Font(bold=True, color="FFFFFF")
    g_cur.fill = grand_fill
    g_cur.alignment = _CENTER
    for col in range(1, 7):
        ws.cell(row=row, column=col).border = _BORDER
    ws.row_dimensions[row].height = 24

    ws.freeze_panes = "A3"
    os.makedirs(output_dir, exist_ok=True)
    filepath = os.path.join(output_dir, f"worker_{year}_{month:02d}.xlsx")
    wb.save(filepath)
    logger.info(f"Worker Excel saved: {filepath}")
    return filepath
